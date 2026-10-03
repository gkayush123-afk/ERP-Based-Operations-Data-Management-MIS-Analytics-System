import json
from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import joinedload

from flask import abort, current_app, flash, make_response, redirect, render_template, request, url_for
from flask_login import current_user

from ..extensions import db
from ..models import AuditLog, Department, User, utcnow
from ..utils.audit import log_action
from ..utils.decorators import roles_required
from . import users_bp
from .forms import ROLE_CHOICES, TemporaryPasswordForm, UserCreateForm, UserEditForm

ROLE_VALUES = {role for role, _label in ROLE_CHOICES}
STATUS_VALUES = {"active", "pending", "rejected", "disabled"}
USERS_PER_PAGE = 20


def active_department_choices(form):
    departments = db.session.scalars(
        db.select(Department).where(Department.is_active.is_(True)).order_by(Department.name)
    ).all()
    form.department_id.choices = [(department.id, department.name) for department in departments]
    return departments


def flash_form_errors(form):
    for errors in form.errors.values():
        for message in errors:
            flash(message, "danger")


@users_bp.route("/", methods=["GET"])
@roles_required("admin")
def index():
    search = request.args.get("q", "").strip()[:120]
    role_filter = request.args.get("role", "")
    status_filter = request.args.get("status", "")
    page = request.args.get("page", 1, type=int)
    page = max(page or 1, 1)
    pending_page = max(request.args.get("pending_page", 1, type=int) or 1, 1)

    query = db.select(User).order_by(User.created_at.desc(), User.id.desc())
    if search:
        term = f"%{search}%"
        query = query.where(
            or_(
                User.username.ilike(term),
                User.full_name.ilike(term),
                User.email.ilike(term),
            )
        )
    if role_filter in ROLE_VALUES:
        query = query.where(User.role == role_filter)
    else:
        role_filter = ""
    if status_filter in STATUS_VALUES:
        query = query.where(User.status == status_filter)
    else:
        status_filter = ""

    pagination = db.paginate(query, page=page, per_page=USERS_PER_PAGE, error_out=False)
    pending_pagination = db.paginate(
        db.select(User)
        .where(User.status == "pending")
        .order_by(User.created_at.asc(), User.id.asc()),
        page=pending_page,
        per_page=USERS_PER_PAGE,
        error_out=False,
    )

    return render_template(
        "users/index.html",
        pagination=pagination,
        pending_pagination=pending_pagination,
        search=search,
        role_filter=role_filter,
        status_filter=status_filter,
        roles=ROLE_CHOICES,
        statuses=sorted(STATUS_VALUES),
    )


@users_bp.route("/create", methods=["GET", "POST"])
@roles_required("admin")
def create():
    form = UserCreateForm()
    try:
        departments = active_department_choices(form)
    except OperationalError:
        db.session.rollback()
        current_app.logger.exception("Database unavailable while loading user creation form.")
        flash(
            "User creation is temporarily unavailable because the database cannot be reached.",
            "danger",
        )
        return render_template("users/create.html", form=form, departments=[]), 503

    if form.validate_on_submit():
        username = form.username.data.strip()
        email = form.email.data.strip().lower()
        try:
            duplicate = db.session.scalar(
                db.select(User.id).where(
                    or_(
                        func.lower(User.username) == username.lower(),
                        func.lower(User.email) == email,
                    )
                )
            )
            if duplicate is not None:
                flash("That username or email is already in use.", "danger")
                return render_template("users/create.html", form=form, departments=departments)

            user = User(
                full_name=form.full_name.data.strip(),
                username=username,
                email=email,
                department_id=form.department_id.data,
                role=form.role.data,
                status="active",
                must_change_password=True,
            )
            user.set_password(form.temporary_password.data)
            db.session.add(user)
            db.session.flush()
            log_action(
                current_user,
                "user.create",
                "user",
                user.id,
                {
                    "username": user.username,
                    "role": user.role,
                    "department_id": user.department_id,
                    "status": user.status,
                },
            )
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash("That username or email is already in use.", "danger")
            return render_template("users/create.html", form=form, departments=departments)
        except OperationalError:
            db.session.rollback()
            current_app.logger.exception("Database unavailable while creating a user.")
            flash(
                "User creation could not be completed because the database cannot be reached.",
                "danger",
            )
            return render_template("users/create.html", form=form, departments=departments), 503

        flash(f"User {user.username} was created. Provide the temporary password securely.", "success")
        return redirect(url_for("users.index"))

    if request.method == "POST":
        flash_form_errors(form)
    return render_template("users/create.html", form=form, departments=departments)


@users_bp.route("/<int:user_id>/edit", methods=["GET", "POST"])
@roles_required("admin")
def edit(user_id):
    user = db.get_or_404(User, user_id)
    form = UserEditForm(obj=user)
    departments = active_department_choices(form)

    if form.validate_on_submit():
        if user.id == current_user.id and form.role.data != "admin":
            flash("You cannot demote your own administrator account.", "danger")
            return redirect(url_for("users.edit", user_id=user.id))

        old_values = {"role": user.role, "department_id": user.department_id}
        user.role = form.role.data
        user.department_id = form.department_id.data
        log_action(
            current_user,
            "user.update",
            "user",
            user.id,
            {
                "before": old_values,
                "after": {"role": user.role, "department_id": user.department_id},
            },
        )
        db.session.commit()
        flash(f"User {user.username} was updated.", "success")
        return redirect(url_for("users.index"))

    if request.method == "POST":
        flash_form_errors(form)
    return render_template("users/edit.html", form=form, user=user, departments=departments)


@users_bp.route("/<int:user_id>/status", methods=["POST"])
@roles_required("admin")
def change_status(user_id):
    user = db.get_or_404(User, user_id)
    requested_status = request.form.get("status")
    if requested_status not in {"active", "disabled"}:
        abort(400)
    if user.id == current_user.id and requested_status == "disabled":
        flash("You cannot deactivate your own administrator account.", "danger")
        return redirect(url_for("users.index"))
    if user.status not in {"active", "disabled"}:
        flash("Only active or deactivated accounts can be changed here.", "warning")
        return redirect(url_for("users.index"))
    if user.status == requested_status:
        flash(f"User {user.username} is already {requested_status}.", "info")
        return redirect(url_for("users.index"))

    previous_status = user.status
    user.status = requested_status
    log_action(
        current_user,
        f"user.{requested_status}",
        "user",
        user.id,
        {"before": {"status": previous_status}, "after": {"status": requested_status}},
    )
    db.session.commit()
    flash(f"User {user.username} was {requested_status}.", "success")
    return redirect(url_for("users.index"))


@users_bp.route("/<int:user_id>/reset-password", methods=["GET", "POST"])
@roles_required("admin")
def reset_password(user_id):
    user = db.get_or_404(User, user_id)
    form = TemporaryPasswordForm()
    if form.validate_on_submit():
        user.set_password(form.temporary_password.data)
        user.must_change_password = True
        user.failed_attempts = 0
        user.locked_until = None
        log_action(
            current_user,
            "user.password_reset",
            "user",
            user.id,
            {"must_change_password": True},
        )
        db.session.commit()
        flash(
            f"Password reset for {user.username}. Share the temporary password securely.",
            "success",
        )
        return redirect(url_for("users.index"))

    if request.method == "POST":
        flash_form_errors(form)
    return render_template("users/reset_password.html", form=form, user=user)


@users_bp.route("/audit-logs", methods=["GET"])
@roles_required("admin")
def audit_logs():
    filters = _audit_filters()
    query = _filtered_audit_query(filters)
    page = max(request.args.get("page", 1, type=int) or 1, 1)
    pagination = db.paginate(
        query.options(joinedload(AuditLog.actor)).order_by(
            AuditLog.occurred_at.desc(), AuditLog.id.desc()
        ),
        page=page,
        per_page=30,
        error_out=False,
    )
    users = db.session.scalars(db.select(User).order_by(User.username)).all()
    actions = db.session.scalars(
        db.select(AuditLog.action).distinct().order_by(AuditLog.action)
    ).all()
    return render_template(
        "users/audit_logs.html",
        pagination=pagination,
        filters=filters,
        users=users,
        actions=actions,
    )


def _audit_filters():
    user_raw = request.args.get("user_id", "").strip()
    try:
        user_id = int(user_raw) if user_raw else None
        if user_id is not None and user_id < 1:
            raise ValueError
    except ValueError:
        abort(400, description="Choose a valid audit-log user.")

    start_raw = request.args.get("start_date", "").strip()
    end_raw = request.args.get("end_date", "").strip()
    try:
        start_date = date.fromisoformat(start_raw) if start_raw else None
        end_date = date.fromisoformat(end_raw) if end_raw else None
    except ValueError:
        abort(400, description="Enter valid audit-log dates.")
    if start_date and end_date and start_date > end_date:
        abort(400, description="Start date must not be after end date.")

    action = request.args.get("action", "").strip()[:80]
    return {
        "user_id": user_id,
        "action": action,
        "start_date": start_date,
        "end_date": end_date,
    }


def _filtered_audit_query(filters):
    query = db.select(AuditLog)
    if filters["user_id"] is not None:
        query = query.where(AuditLog.actor_user_id == filters["user_id"])
    if filters["action"]:
        query = query.where(AuditLog.action == filters["action"])
    if filters["start_date"]:
        query = query.where(
            AuditLog.occurred_at >= datetime.combine(filters["start_date"], time.min)
        )
    if filters["end_date"] and filters["end_date"] < date.max:
        query = query.where(
            AuditLog.occurred_at
            < datetime.combine(filters["end_date"] + timedelta(days=1), time.min)
        )
    return query


def _safe_excel_text(value):
    text = str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


@users_bp.route("/audit-logs/export", methods=["GET"])
@roles_required("admin")
def export_audit_logs():
    filters = _audit_filters()
    query = _filtered_audit_query(filters).options(joinedload(AuditLog.actor)).order_by(
        AuditLog.occurred_at.desc(), AuditLog.id.desc()
    )
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Audit Logs")
    row_count = db.session.scalar(
        _filtered_audit_query(filters).with_only_columns(func.count(AuditLog.id))
    ) or 0
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{row_count + 1}"
    for letter, width in {
        "A": 22,
        "B": 24,
        "C": 30,
        "D": 20,
        "E": 14,
        "F": 20,
        "G": 60,
    }.items():
        sheet.column_dimensions[letter].width = width
    headers = ("Time (UTC)", "User", "Action", "Entity", "Entity ID", "IP address", "Details")
    header = []
    for label in headers:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A7E")
        header.append(cell)
    sheet.append(header)
    for entry in db.session.scalars(query).yield_per(500):
        username = entry.actor.username if entry.actor else entry.details.get("username", "Unknown")
        sheet.append(
            [
                entry.occurred_at.isoformat(sep=" "),
                _safe_excel_text(username),
                _safe_excel_text(entry.action),
                _safe_excel_text(entry.entity),
                entry.entity_id,
                entry.ip_address or "",
                _safe_excel_text(json.dumps(entry.details, sort_keys=True, ensure_ascii=True)),
            ]
        )
    output = BytesIO()
    workbook.save(output)

    log_action(
        current_user,
        "audit.export",
        "audit_log",
        None,
        {
            "user_id": filters["user_id"],
            "action": filters["action"],
            "start_date": filters["start_date"].isoformat() if filters["start_date"] else None,
            "end_date": filters["end_date"].isoformat() if filters["end_date"] else None,
        },
    )
    db.session.commit()
    response = make_response(output.getvalue())
    response.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = 'attachment; filename="audit_logs.xlsx"'
    return response


@users_bp.route("/<int:user_id>/approval/<decision>", methods=["POST"])
@roles_required("admin")
def decide_registration(user_id, decision):
    if decision not in {"approve", "reject"}:
        abort(404)
    user = db.get_or_404(User, user_id)
    if user.status != "pending":
        flash("That registration is no longer pending.", "warning")
        return redirect(url_for("users.index"))

    user.status = "active" if decision == "approve" else "rejected"
    user.approved_by = current_user.id
    user.approved_at = utcnow()
    log_action(
        current_user,
        f"user.{decision}",
        "user",
        user.id,
        {"username": user.username, "status": user.status},
    )
    db.session.commit()
    result = "approved" if decision == "approve" else "rejected"
    flash(f"Registration for {user.username} was {result}.", "success")
    return redirect(url_for("users.index"))
