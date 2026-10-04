from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from email_validator import EmailNotValidError, validate_email
from flask import abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError
from werkzeug.utils import secure_filename

from ..extensions import db
from ..models import AttendanceRecord, Department, Employee, OperationRecord, utcnow
from ..utils.audit import log_action
from ..utils.decorators import roles_required
from . import employees_bp
from .forms import EmployeeForm, STATUS_CHOICES, valid_phone

PER_PAGE = 20
IMPORT_COLUMNS = (
    "employee_code",
    "full_name",
    "email",
    "phone",
    "designation",
    "department",
    "joining_date",
    "status",
)
SORT_COLUMNS = {
    "employee_code": Employee.employee_code,
    "full_name": Employee.full_name,
    "email": Employee.email,
    "department": Department.name,
    "joining_date": Employee.joining_date,
    "status": Employee.status,
    "created_at": Employee.created_at,
}


def get_departments():
    query = db.select(Department).where(Department.is_active.is_(True)).order_by(Department.name)
    if current_user.role in {"manager", "data_entry"}:
        query = query.where(Department.id == current_user.department_id)
    return db.session.scalars(query).all()


def set_department_choices(form, include_department=None):
    departments = get_departments()
    if (
        include_department is not None
        and all(department.id != include_department.id for department in departments)
    ):
        departments.append(include_department)
        departments.sort(key=lambda department: department.name.casefold())
    form.department_id.choices = [(department.id, department.name) for department in departments]
    if current_user.role == "data_entry" and current_user.department_id:
        form.department_id.data = current_user.department_id
    return departments


def employee_query():
    query = db.select(Employee).join(Department).where(Employee.deleted_at.is_(None))
    if current_user.role in {"manager", "data_entry"}:
        query = query.where(Employee.department_id == current_user.department_id)
    return query


def normalize_code(value):
    return str(value or "").strip()


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError):
        return None


@employees_bp.route("/")
@roles_required("admin", "manager", "data_entry")
def index():
    search = request.args.get("q", "").strip()[:120]
    status_filter = request.args.get("status", "")
    department_filter = request.args.get("department_id", type=int)
    sort_key = request.args.get("sort", "employee_code")
    direction = request.args.get("direction", "asc")
    page = max(request.args.get("page", 1, type=int) or 1, 1)
    if sort_key not in SORT_COLUMNS:
        sort_key = "employee_code"
    if direction not in {"asc", "desc"}:
        direction = "asc"

    query = employee_query()
    if search:
        term = f"%{search}%"
        query = query.where(
            or_(
                Employee.employee_code.ilike(term),
                Employee.full_name.ilike(term),
                Employee.email.ilike(term),
                Employee.designation.ilike(term),
            )
        )
    if status_filter in {"active", "inactive"}:
        query = query.where(Employee.status == status_filter)
    else:
        status_filter = ""
    if department_filter is not None:
        if current_user.role == "data_entry":
            department_filter = current_user.department_id
        if current_user.role == "manager" and department_filter != current_user.department_id:
            abort(403)
        query = query.where(Employee.department_id == department_filter)

    sort_column = SORT_COLUMNS[sort_key]
    query = query.order_by(
        sort_column.desc() if direction == "desc" else sort_column.asc(),
        Employee.id.asc(),
    )
    pagination = db.paginate(query, page=page, per_page=PER_PAGE, error_out=False)
    return render_template(
        "employees/index.html",
        pagination=pagination,
        departments=get_departments(),
        search=search,
        status_filter=status_filter,
        department_filter=department_filter,
        sort_key=sort_key,
        direction=direction,
        statuses=STATUS_CHOICES,
        active_nav="employees",
    )


@employees_bp.route("/<int:employee_id>")
@roles_required("admin", "manager", "data_entry")
def show(employee_id):
    employee = db.get_or_404(Employee, employee_id)
    if employee.deleted_at is not None:
        abort(404)
    if current_user.role in {"manager", "data_entry"}:
        if employee.department_id != current_user.department_id:
            abort(403)
    recent_attendance = db.session.scalars(
        db.select(AttendanceRecord)
        .where(AttendanceRecord.employee_id == employee.id)
        .order_by(AttendanceRecord.attendance_date.desc(), AttendanceRecord.id.desc())
        .limit(5)
    ).all()
    operation_counts = {
        status: 0 for status in ("verified", "pending", "rejected")
    }
    for status, total in db.session.execute(
        db.select(OperationRecord.status, func.count(OperationRecord.id))
        .where(OperationRecord.employee_id == employee.id)
        .group_by(OperationRecord.status)
    ):
        if status in operation_counts:
            operation_counts[status] = int(total)
    operation_counts["total"] = sum(operation_counts.values())
    return render_template(
        "employees/show.html",
        employee=employee,
        recent_attendance=recent_attendance,
        operation_counts=operation_counts,
        active_nav="employees",
    )


@employees_bp.route("/create", methods=["GET", "POST"])
@roles_required("admin", "manager", "data_entry")
def create():
    form = EmployeeForm()
    departments = set_department_choices(form)
    if not departments:
        flash("No active department is available for your account.", "danger")
        return redirect(url_for("employees.index"))

    if form.validate_on_submit():
        code = form.employee_code.data.strip()
        if db.session.scalar(
            db.select(Employee.id).where(func.lower(Employee.employee_code) == code.lower())
        ):
            flash("That employee code is already in use.", "danger")
            return render_template("employees/form.html", form=form, employee=None, departments=departments, active_nav="employees")

        department_id = (
            current_user.department_id
            if current_user.role == "data_entry"
            else form.department_id.data
        )
        if not db.session.get(Department, department_id):
            abort(400)
        employee = Employee(
            employee_code=code,
            full_name=form.full_name.data.strip(),
            email=form.email.data.strip().lower(),
            phone=(form.phone.data or "").strip() or None,
            designation=form.designation.data.strip(),
            department_id=department_id,
            joining_date=form.joining_date.data,
            status=form.status.data,
            created_by=current_user.id,
        )
        db.session.add(employee)
        try:
            db.session.flush()
            log_action(
                current_user,
                "employee.create",
                "employee",
                employee.id,
                {
                    "employee_code": employee.employee_code,
                    "department_id": employee.department_id,
                    "status": employee.status,
                },
            )
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash("That employee code is already in use.", "danger")
            return render_template("employees/form.html", form=form, employee=None, departments=departments, active_nav="employees")
        flash(f"Employee {employee.employee_code} was created.", "success")
        return redirect(url_for("employees.index"))

    if request.method == "POST":
        for errors in form.errors.values():
            for message in errors:
                flash(message, "danger")
    return render_template("employees/form.html", form=form, employee=None, departments=departments, active_nav="employees")


@employees_bp.route("/<int:employee_id>/edit", methods=["GET", "POST"])
@roles_required("admin", "manager", "data_entry")
def edit(employee_id):
    employee = db.get_or_404(Employee, employee_id)
    if employee.deleted_at is not None:
        abort(404)
    if current_user.role == "data_entry":
        if employee.department_id != current_user.department_id:
            abort(403)
        if employee.created_by != current_user.id:
            abort(403)
    if current_user.role == "manager" and employee.department_id != current_user.department_id:
        abort(403)

    form = EmployeeForm(obj=employee)
    departments = set_department_choices(form, employee.department)
    if form.validate_on_submit():
        code = form.employee_code.data.strip()
        duplicate = db.session.scalar(
            db.select(Employee.id).where(
                func.lower(Employee.employee_code) == code.lower(),
                Employee.id != employee.id,
            )
        )
        if duplicate is not None:
            flash("That employee code is already in use.", "danger")
            return render_template("employees/form.html", form=form, employee=employee, departments=departments, active_nav="employees")

        old = {
            "employee_code": employee.employee_code,
            "full_name": employee.full_name,
            "email": employee.email,
            "phone": employee.phone,
            "designation": employee.designation,
            "department_id": employee.department_id,
            "joining_date": employee.joining_date.isoformat(),
            "status": employee.status,
        }
        employee.employee_code = code
        employee.full_name = form.full_name.data.strip()
        employee.email = form.email.data.strip().lower()
        employee.phone = (form.phone.data or "").strip() or None
        employee.designation = form.designation.data.strip()
        if current_user.role != "data_entry":
            employee.department_id = form.department_id.data
        employee.joining_date = form.joining_date.data
        employee.status = form.status.data
        log_action(
            current_user,
            "employee.update",
            "employee",
            employee.id,
            {
                "before": old,
                "after": {
                    "employee_code": employee.employee_code,
                    "full_name": employee.full_name,
                    "department_id": employee.department_id,
                    "status": employee.status,
                },
            },
        )
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash("That employee code is already in use.", "danger")
            return render_template("employees/form.html", form=form, employee=employee, departments=departments, active_nav="employees")
        flash(f"Employee {employee.employee_code} was updated.", "success")
        return redirect(url_for("employees.index"))

    if request.method == "POST":
        for errors in form.errors.values():
            for message in errors:
                flash(message, "danger")
    return render_template("employees/form.html", form=form, employee=employee, departments=departments, active_nav="employees")


@employees_bp.route("/<int:employee_id>/delete", methods=["POST"])
@roles_required("admin")
def delete(employee_id):
    employee = db.get_or_404(Employee, employee_id)
    if employee.deleted_at is not None:
        flash("This employee is already deleted.", "info")
        return redirect(url_for("employees.index"))

    employee.deleted_at = utcnow()
    employee.deleted_by = current_user.id
    log_action(
        current_user,
        "employee.soft_delete",
        "employee",
        employee.id,
        {"employee_code": employee.employee_code, "full_name": employee.full_name},
    )
    db.session.commit()
    flash(f"Employee {employee.employee_code} was deleted.", "success")
    return redirect(url_for("employees.index"))


@employees_bp.route("/import", methods=["GET", "POST"])
@roles_required("admin", "manager", "data_entry")
def import_excel():
    departments = get_departments()
    row_results = []
    if request.method == "POST":
        upload = request.files.get("file")
        if upload is None or not upload.filename:
            flash("Choose an .xlsx file to import.", "danger")
            return render_template("employees/import.html", row_results=row_results, departments=departments, active_nav="employees")
        if Path(secure_filename(upload.filename)).suffix.lower() != ".xlsx":
            flash("Only .xlsx workbooks are supported.", "danger")
            return render_template("employees/import.html", row_results=row_results, departments=departments, active_nav="employees")

        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            headers = [str(value or "").strip().lower() for value in (header_row or ())]
            missing = [column for column in IMPORT_COLUMNS if column not in headers]
            if missing:
                workbook.close()
                flash(f"Missing required column(s): {', '.join(missing)}.", "danger")
                return render_template("employees/import.html", row_results=row_results, departments=departments, active_nav="employees")
        except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError, ParseError):
            flash("The uploaded workbook could not be read. Please select a valid .xlsx file.", "danger")
            return render_template("employees/import.html", row_results=row_results, departments=departments, active_nav="employees")

        column_index = {name: headers.index(name) for name in IMPORT_COLUMNS}
        seen_codes = set()
        for row_number, values in enumerate(rows, start=2):
            if not values or not any(value not in (None, "") for value in values):
                continue
            row = {
                name: values[index] if index < len(values) else None
                for name, index in column_index.items()
            }
            errors = []
            code = normalize_code(row["employee_code"])
            full_name = str(row["full_name"] or "").strip()
            email = str(row["email"] or "").strip().lower()
            phone = str(row["phone"] or "").strip()
            designation = str(row["designation"] or "").strip()
            if not code:
                errors.append("employee_code is required")
            elif len(code) > 40:
                errors.append("employee_code must be at most 40 characters")
            elif code.lower() in seen_codes:
                errors.append("employee_code is duplicated in this workbook")
            elif db.session.scalar(
                db.select(Employee.id).where(func.lower(Employee.employee_code) == code.lower())
            ):
                errors.append("employee_code already exists")
            if not full_name:
                errors.append("full_name is required")
            if not email:
                errors.append("email is required")
            else:
                try:
                    validate_email(email, check_deliverability=False)
                except EmailNotValidError:
                    errors.append("email is invalid")
            if not valid_phone(phone):
                errors.append("phone format is invalid")
            if not designation:
                errors.append("designation is required")
            joining_date = normalize_date(row["joining_date"])
            if joining_date is None:
                errors.append("joining_date must be a valid date (YYYY-MM-DD)")
            elif joining_date > date.today():
                errors.append("joining_date cannot be in the future")
            status = str(row["status"] or "").strip().lower()
            if status not in {"active", "inactive"}:
                errors.append("status must be active or inactive")

            department_id = current_user.department_id
            if current_user.role != "data_entry":
                department_value = str(row["department"] or "").strip()
                department = None
                if department_value.isdigit():
                    department = db.session.get(Department, int(department_value))
                    if department is not None and not department.is_active:
                        department = None
                if department is None:
                    department = db.session.scalar(
                        db.select(Department).where(
                            func.lower(Department.name) == department_value.lower(),
                            Department.is_active.is_(True),
                        )
                    )
                if department is not None and current_user.role == "manager":
                    if department.id != current_user.department_id:
                        department = None
                if department is None:
                    if current_user.role == "manager":
                        errors.append("department must be your own active department")
                    else:
                        errors.append("department does not match an active department")
                else:
                    department_id = department.id
            elif not department_id:
                errors.append("your account has no department assigned")

            if errors:
                row_results.append(
                    {"row": row_number, "status": "error", "message": "; ".join(errors)}
                )
                if code:
                    seen_codes.add(code.lower())
                continue

            employee = Employee(
                employee_code=code,
                full_name=full_name,
                email=email,
                phone=phone or None,
                designation=designation,
                department_id=department_id,
                joining_date=joining_date,
                status=status,
                created_by=current_user.id,
            )
            db.session.add(employee)
            try:
                db.session.flush()
                log_action(
                    current_user,
                    "employee.import",
                    "employee",
                    employee.id,
                    {
                        "employee_code": employee.employee_code,
                        "department_id": employee.department_id,
                        "source_row": row_number,
                    },
                )
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                row_results.append(
                    {"row": row_number, "status": "error", "message": "employee_code already exists"}
                )
                seen_codes.add(code.lower())
            else:
                row_results.append(
                    {"row": row_number, "status": "success", "message": f"Imported {code}"}
                )
                seen_codes.add(code.lower())
        workbook.close()
        imported_count = sum(result["status"] == "success" for result in row_results)
        error_count = len(row_results) - imported_count
        flash(f"Import complete: {imported_count} imported, {error_count} row(s) failed.", "info")
    return render_template("employees/import.html", row_results=row_results, departments=departments, active_nav="employees")
