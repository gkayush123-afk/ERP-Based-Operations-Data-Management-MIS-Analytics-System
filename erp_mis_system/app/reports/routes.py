from datetime import date
from io import BytesIO

from flask import abort, make_response, render_template, request
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..extensions import db
from ..models import Department, Employee, utcnow
from ..utils.audit import log_action
from ..utils.decorators import roles_required
from . import reports_bp
from .service import (
    STATUS_LABELS,
    STATUS_VALUES,
    build_operation_by_department,
    build_operation_by_employee,
    build_operation_summary,
    build_operation_trend,
    build_report,
    pct_change,
    previous_period,
    report_period,
)


def _filters():
    report_type = request.args.get("type", "daily")
    today = date.today()
    try:
        period_date = date.fromisoformat(request.args.get("date", today.isoformat()))
        month_raw = request.args.get("month", today.strftime("%Y-%m"))
        period_month = date.fromisoformat(f"{month_raw}-01")
        start_date = date.fromisoformat(request.args.get("start_date", today.replace(day=1).isoformat()))
        end_date = date.fromisoformat(request.args.get("end_date", today.isoformat()))
    except ValueError:
        abort(400, description="Enter valid report dates.")
    try:
        start_date, end_date = report_period(
            report_type,
            period_date=period_date,
            period_month=period_month,
            start_date=start_date,
            end_date=end_date,
        )
    except ValueError as error:
        abort(400, description=str(error))

    status = request.args.get("status") or None
    if status not in (*STATUS_VALUES, None):
        abort(400, description="Choose a valid attendance status.")
    department_raw = request.args.get("department_id", "")
    try:
        department_id = int(department_raw) if department_raw else None
        if department_id is not None and department_id < 1:
            raise ValueError
    except ValueError:
        abort(400, description="Choose a valid department.")
    if current_user.role == "manager":
        if current_user.department_id is None:
            abort(403, description="A manager must be assigned to a department.")
        department_id = current_user.department_id
    return {
        "type": report_type,
        "period_date": period_date,
        "period_month": period_month,
        "start_date": start_date,
        "end_date": end_date,
        "department_id": department_id,
        "status": status,
    }


def _report_data(filters):
    return build_report(
        db.session,
        filters["type"],
        filters["start_date"],
        filters["end_date"],
        department_id=filters["department_id"],
        status=filters["status"],
    )


def _department_options(filters):
    query = db.select(Department).where(Department.is_active.is_(True))
    if current_user.role == "manager":
        query = query.where(Department.id == current_user.department_id)
    return db.session.scalars(query.order_by(Department.name)).all()


def _chart_payload(rows, report_type):
    if report_type == "summary":
        departments = sorted({row["department"] for row in rows})
        return {
            "labels": departments,
            "series": {
                status: [
                    sum(row[status] for row in rows if row["department"] == name)
                    for name in departments
                ]
                for status in STATUS_VALUES
            },
        }
    if report_type == "department":
        return {
            "labels": [row["department"] for row in rows],
            "series": {
                status: [row[status] for row in rows]
                for status in STATUS_VALUES
            },
        }
    totals = {}
    for row in rows:
        date_label = row["date"]
        values = totals.setdefault(date_label, {status: 0 for status in STATUS_VALUES})
        for status in STATUS_VALUES:
            values[status] += row[status]
    return {
        "labels": list(totals),
        "series": {
            status: [totals[label][status] for label in totals]
            for status in STATUS_VALUES
        },
    }


def _excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


@reports_bp.route("/")
@roles_required("admin", "manager")
def index():
    filters = _filters()
    rows = _report_data(filters)
    departments = _department_options(filters)
    return render_template(
        "reports/index.html",
        rows=rows,
        departments=departments,
        filters=filters,
        chart=_chart_payload(rows, filters["type"]),
        status_labels=STATUS_LABELS,
    )


@reports_bp.route("/export/<file_format>")
@roles_required("admin", "manager")
def export(file_format):
    if file_format not in {"xlsx", "pdf"}:
        abort(404)
    filters = _filters()
    rows = _report_data(filters)
    if filters["type"] == "summary":
        columns = [
            ("department", "Department"),
            ("employee_code", "Employee ID"),
            ("employee_name", "Employee Name"),
            ("present", "Present"),
            ("absent", "Absent"),
            ("late", "Late"),
            ("leave", "Leave"),
            ("half_day", "Half-day"),
            ("total", "Total"),
            ("attendance_percentage", "Attendance %"),
        ]
    else:
        columns = (
            [
                ("department", "Department"),
                ("present", "Present"),
                ("absent", "Absent"),
                ("late", "Late"),
                ("leave", "Leave"),
                ("half_day", "Half-day"),
                ("active_employees", "Active employees"),
                ("attendance_percentage", "Attendance %"),
            ]
            if filters["type"] == "department"
            else [
                ("date", "Date"),
                ("department", "Department"),
                ("present", "Present"),
                ("absent", "Absent"),
                ("late", "Late"),
                ("leave", "Leave"),
                ("half_day", "Half-day"),
            ]
        )
    report_title = f"{filters['type'].title()} MIS Report"
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MIS Report"
        sheet.append([report_title])
        sheet.append(
            [
                f"Period: {filters['start_date'].isoformat()} to {filters['end_date'].isoformat()}",
                f"Status filter: {STATUS_LABELS.get(filters['status'], 'All')}",
            ]
        )
        sheet.append([label for _, label in columns])
        for cell in sheet[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174A7E")
        for row in rows:
            sheet.append([_excel_value(row[key]) for key, _ in columns])
        for column_cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 36)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        output = BytesIO()
        workbook.save(output)
        payload = output.getvalue()
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"{filters['type']}_mis_report.xlsx"
    else:
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        data = [[label for _, label in columns]]
        data.extend([[str(row[key]) for key, _ in columns] for row in rows])
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#174a7e")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccd5df")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f6fa")]),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        document.build(
            [
                Paragraph(report_title, styles["Title"]),
                Paragraph(
                    f"Period: {filters['start_date'].isoformat()} to {filters['end_date'].isoformat()}",
                    styles["Normal"],
                ),
                Spacer(1, 12),
                table,
            ]
        )
        payload = output.getvalue()
        content_type = "application/pdf"
        filename = f"{filters['type']}_mis_report.pdf"

    log_action(
        current_user,
        "reports.export",
        "mis_report",
        None,
        {
            "report_type": filters["type"],
            "format": file_format,
            "start_date": filters["start_date"].isoformat(),
            "end_date": filters["end_date"].isoformat(),
            "department_id": filters["department_id"],
            "status": filters["status"],
        },
    )
    db.session.commit()
    response = make_response(payload)
    response.headers["Content-Type"] = content_type
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


OPERATION_REPORT_TYPES = (
    ("overall", "Overall Summary"),
    ("department", "Department-wise"),
    ("employee", "Employee-wise"),
)
OPERATION_TYPE_LABELS = dict(OPERATION_REPORT_TYPES)
MAX_OPERATION_RANGE_DAYS = 366


def _operation_filters():
    report_type = request.args.get("type", "overall")
    if report_type not in OPERATION_TYPE_LABELS:
        abort(400, description="Choose a valid report type.")
    today = date.today()
    try:
        start_date = date.fromisoformat(
            request.args.get("start_date", today.replace(day=1).isoformat())
        )
        end_date = date.fromisoformat(request.args.get("end_date", today.isoformat()))
    except ValueError:
        abort(400, description="Enter valid report dates.")
    if start_date > end_date:
        abort(400, description="Start date must not be after end date.")
    if (end_date - start_date).days > MAX_OPERATION_RANGE_DAYS:
        abort(400, description="Date range must not exceed one year.")
    department_id = request.args.get("department_id", type=int)
    employee_id = request.args.get("employee_id", type=int)
    search = (request.args.get("q") or "").strip()[:80]
    if current_user.role == "manager":
        if current_user.department_id is None:
            abort(403, description="A manager must be assigned to a department.")
        department_id = current_user.department_id
    elif department_id is not None:
        department = db.session.get(Department, department_id)
        if department is None:
            abort(400, description="Choose a valid department.")
    if employee_id is not None:
        employee = db.session.get(Employee, employee_id)
        if employee is None or employee.deleted_at is not None:
            abort(400, description="Choose a valid employee.")
        if (
            current_user.role == "manager"
            and employee.department_id != current_user.department_id
        ):
            abort(403)
    return {
        "type": report_type,
        "start_date": start_date,
        "end_date": end_date,
        "department_id": department_id,
        "employee_id": employee_id,
        "q": search,
    }


def _operation_options(filters):
    department_query = db.select(Department).where(Department.is_active.is_(True))
    if current_user.role == "manager":
        department_query = department_query.where(
            Department.id == current_user.department_id
        )
    departments = db.session.scalars(department_query.order_by(Department.name)).all()
    employee_query = db.select(Employee).where(
        Employee.deleted_at.is_(None), Employee.status == "active"
    )
    if current_user.role == "manager":
        employee_query = employee_query.where(
            Employee.department_id == current_user.department_id
        )
    elif filters["department_id"] is not None:
        employee_query = employee_query.where(
            Employee.department_id == filters["department_id"]
        )
    employees = db.session.scalars(
        employee_query.order_by(Employee.employee_code).limit(500)
    ).all()
    return departments, employees


def _operation_cards(session, filters):
    current = build_operation_summary(
        session,
        filters["start_date"],
        filters["end_date"],
        department_id=filters["department_id"],
        employee_id=filters["employee_id"],
        search=filters["q"],
    )
    prev_start, prev_end = previous_period(filters["start_date"], filters["end_date"])
    previous = build_operation_summary(
        session,
        prev_start,
        prev_end,
        department_id=filters["department_id"],
        employee_id=filters["employee_id"],
        search=filters["q"],
    )
    specs = (
        ("total", "Total Records", True),
        ("verified", "Verified Records", True),
        ("pending", "Pending Records", False),
        ("rejected", "Rejected Records", False),
    )
    cards = []
    for key, label, up_is_good in specs:
        percent, direction = pct_change(current[key], previous[key])
        good = direction == "flat" or (direction == "up") == up_is_good
        cards.append(
            {
                "key": key,
                "label": label,
                "value": current[key],
                "percent": percent,
                "direction": direction,
                "good": good,
            }
        )
    return cards


@reports_bp.route("/operations")
@roles_required("admin", "manager")
def operations():
    filters = _operation_filters()
    cards = _operation_cards(db.session, filters)
    trend = build_operation_trend(
        db.session,
        filters["start_date"],
        filters["end_date"],
        department_id=filters["department_id"],
        employee_id=filters["employee_id"],
        search=filters["q"],
    )
    dept_rows = build_operation_by_department(
        db.session,
        filters["start_date"],
        filters["end_date"],
        department_id=filters["department_id"],
        employee_id=filters["employee_id"],
        search=filters["q"],
    )
    emp_rows = build_operation_by_employee(
        db.session,
        filters["start_date"],
        filters["end_date"],
        department_id=filters["department_id"],
        employee_id=filters["employee_id"],
        search=filters["q"],
    )
    top_departments = dept_rows[:5]
    others_total = sum(row["total"] for row in dept_rows[5:])
    donut_labels = [row["department"] for row in top_departments]
    donut_values = [row["total"] for row in top_departments]
    if others_total:
        donut_labels.append("Others")
        donut_values.append(others_total)
    donut_legend = [
        {
            "label": label,
            "value": value,
            "percent": round(value * 100 / sum(donut_values), 1) if sum(donut_values) else 0.0,
        }
        for label, value in zip(donut_labels, donut_values)
    ]
    top_employees = emp_rows[:5]
    departments, employees = _operation_options(filters)
    dept_id_by_name = {item.name: item.id for item in departments}
    show_full = request.args.get("full") == "1"
    table_rows = dept_rows if filters["type"] != "employee" else emp_rows
    if not show_full:
        table_rows = table_rows[:10]
    return render_template(
        "reports/operations.html",
        active_nav="reports",
        filters=filters,
        report_types=OPERATION_REPORT_TYPES,
        type_label=OPERATION_TYPE_LABELS[filters["type"]],
        cards=cards,
        trend=trend,
        donut_labels=donut_labels,
        donut_values=donut_values,
        donut_legend=donut_legend,
        donut_total=sum(donut_values),
        top_employees=top_employees,
        table_rows=table_rows,
        show_full=show_full,
        departments=departments,
        dept_id_by_name=dept_id_by_name,
        employees=employees,
    )


@reports_bp.route("/operations/export/<file_format>")
@roles_required("admin", "manager")
def operations_export(file_format):
    if file_format not in {"xlsx", "pdf"}:
        abort(404)
    filters = _operation_filters()
    cards = _operation_cards(db.session, filters)
    if filters["type"] == "employee":
        table_rows = build_operation_by_employee(
            db.session,
            filters["start_date"],
            filters["end_date"],
            department_id=filters["department_id"],
            employee_id=filters["employee_id"],
            search=filters["q"],
        )
        columns = [
            ("department", "Department"),
            ("employee_code", "Employee ID"),
            ("employee_name", "Employee Name"),
            ("total", "Total"),
            ("verified", "Verified"),
            ("pending", "Pending"),
            ("rejected", "Rejected"),
            ("needs_correction", "Needs Correction"),
            ("completion_pct", "Completion %"),
            ("last_updated", "Last Updated"),
        ]
    else:
        table_rows = build_operation_by_department(
            db.session,
            filters["start_date"],
            filters["end_date"],
            department_id=filters["department_id"],
            employee_id=filters["employee_id"],
            search=filters["q"],
        )
        columns = [
            ("department", "Department"),
            ("total", "Total"),
            ("verified", "Verified"),
            ("pending", "Pending"),
            ("rejected", "Rejected"),
            ("needs_correction", "Needs Correction"),
            ("completion_pct", "Completion %"),
            ("last_updated", "Last Updated"),
        ]
    generated_by = current_user.full_name
    generated_at = utcnow().isoformat(sep=" ", timespec="minutes")
    title = f"MIS Reports - {OPERATION_TYPE_LABELS[filters['type']]}"
    period = f"{filters['start_date'].isoformat()} to {filters['end_date'].isoformat()}"
    card_rows = [
        [card["label"], card["value"], f"{card['percent']}% vs previous period"]
        for card in cards
    ]
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MIS Report"
        sheet.append([title])
        sheet.append([f"Period: {period}", f"Generated by {generated_by} at {generated_at}"])
        sheet.append([])
        sheet.append(["Summary", "Value", "Change"])
        for row in card_rows:
            sheet.append([_excel_value(value) for value in row])
        sheet.append([])
        sheet.append([label for _, label in columns])
        for row in table_rows:
            sheet.append([_excel_value(row[key]) for key, _ in columns])
        for cell in sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174A7E")
        header_row = 4 + len(card_rows) + 2
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174A7E")
        for column_cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 36)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        output = BytesIO()
        workbook.save(output)
        payload = output.getvalue()
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"mis_report_{filters['type']}.xlsx"
    else:
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        story = [
            Paragraph(f"<b>{title}</b>", styles["Title"]),
            Paragraph(f"Period: {period}", styles["Normal"]),
            Paragraph(f"Generated by {generated_by} at {generated_at}", styles["Normal"]),
            Spacer(1, 12),
            Paragraph("<b>Summary</b>", styles["Heading3"]),
            Table(
                [["Summary", "Value", "Change"]] + card_rows,
                repeatRows=1,
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#174a7e")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccd5df")),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]
                ),
            ),
            Spacer(1, 12),
            Paragraph("<b>Detailed Report</b>", styles["Heading3"]),
            Table(
                [[label for _, label in columns]]
                + [[str(row[key]) for key, _ in columns] for row in table_rows],
                repeatRows=1,
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#174a7e")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#ccd5df")),
                        ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ]
                ),
            ),
        ]
        document.build(story)
        payload = output.getvalue()
        content_type = "application/pdf"
        filename = f"mis_report_{filters['type']}.pdf"

    log_action(
        current_user,
        "mis_reports.export",
        "mis_report",
        None,
        {
            "report_type": filters["type"],
            "format": file_format,
            "start_date": filters["start_date"].isoformat(),
            "end_date": filters["end_date"].isoformat(),
            "department_id": filters["department_id"],
            "employee_id": filters["employee_id"],
            "generated_by": generated_by,
        },
    )
    db.session.commit()
    response = make_response(payload)
    response.headers["Content-Type"] = content_type
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
